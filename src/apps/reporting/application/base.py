from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path
from typing import Any, ClassVar, override
from uuid import UUID

import polars as pl
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from utils.pagination import SearchRequest


@dataclass(frozen=True, slots=True)
class ReportActor:
    """Stable requester facts available to authorization-scoped report selects."""

    id: UUID
    is_superuser: bool


@dataclass(frozen=True, slots=True)
class ChartData:
    """Small pre-aggregated chart dataset safe to keep in memory."""

    title: str
    categories: tuple[str, ...]
    values: tuple[int, ...]


@dataclass(frozen=True, slots=True)
class ProcessedReport:
    """Transformed tabular data and its bounded presentation summary."""

    title: str
    data_sheet_title: str
    dataframe: pl.DataFrame
    metrics: tuple[tuple[str, int], ...]
    charts: tuple[ChartData, ...]


class BasePolarsReport(ABC):
    """Own a report's ORM selection and Polars transformations."""

    query_type: ClassVar[type[SearchRequest]]

    @classmethod
    @abstractmethod
    def build_statement(cls, query: SearchRequest, actor: ReportActor) -> Any:
        """Build an ORM-only, permission-scoped, deterministically ordered select."""

    @classmethod
    @abstractmethod
    def build_count_statement(cls, query: SearchRequest, actor: ReportActor) -> Any:
        """Build the matching count select using exactly the same filters."""

    @abstractmethod
    def transform_chunk(self, frame: pl.DataFrame, language: str) -> pl.DataFrame:
        """Validate and transform one bounded result partition."""

    @abstractmethod
    def finalize(self, chunks: list[pl.DataFrame], language: str) -> ProcessedReport:
        """Create the final dataframe and bounded summary, consuming all chunks."""


class BaseExcelRenderer(ABC):
    """Render processed report data without owning queries, storage, or task state."""

    @abstractmethod
    def render(self, report: ProcessedReport, destination: Path) -> None:
        """Write an XLSX workbook to *destination*."""


class CleanExcelRenderer(BaseExcelRenderer):
    """Create a colorful, polished, write-only workbook with bounded charts."""

    _navy = "172A3A"
    _blue = "2F75B5"
    _pale_blue = "DCE6F1"
    _teal = "2A9D8F"
    _gold = "E9C46A"
    _coral = "E76F51"
    _white = "FFFFFF"
    _light = "F5F7FA"

    @staticmethod
    def _cell(
        sheet: Any,
        value: Any,
        *,
        bold: bool = False,
        color: str = "172A3A",
        fill: str | None = None,
        size: int = 11,
        align: str = "left",
        number_format: str | None = None,
    ) -> Any:
        cell = WriteOnlyCell(sheet, value=value)
        cell.font = Font(name="Aptos", bold=bold, color=color, size=size)
        if bool(fill):
            cell.fill = PatternFill(fill_type="solid", fgColor=fill)
        cell.alignment = Alignment(horizontal=align, vertical="center")
        if bool(number_format):
            cell.number_format = number_format
        return cell

    @override
    def render(self, report: ProcessedReport, destination: Path) -> None:
        workbook = Workbook(write_only=True)
        try:
            summary = workbook.create_sheet("Summary")
            summary.sheet_properties.tabColor = self._blue
            summary.sheet_view.showGridLines = False
            summary.column_dimensions["A"].width = 30
            summary.column_dimensions["B"].width = 18
            summary.append(
                [
                    self._cell(
                        summary,
                        report.title,
                        bold=True,
                        color=self._white,
                        fill=self._navy,
                        size=18,
                    ),
                    self._cell(summary, "", fill=self._navy),
                ]
            )
            summary.append([self._cell(summary, ""), self._cell(summary, "")])
            summary.append(
                [
                    self._cell(
                        summary,
                        "Metric",
                        bold=True,
                        color=self._white,
                        fill=self._blue,
                    ),
                    self._cell(
                        summary,
                        "Value",
                        bold=True,
                        color=self._white,
                        fill=self._blue,
                        align="right",
                    ),
                ]
            )
            for index, (label, value) in enumerate(report.metrics, start=1):
                palette = (self._pale_blue, "D9EAD3", self._gold, self._coral)
                fill = palette[(index - 1) % len(palette)]
                summary.append(
                    [
                        self._cell(summary, label, fill=fill),
                        self._cell(summary, value, fill=fill, align="right", number_format="#,##0"),
                    ]
                )

            current_row = len(report.metrics) + 3
            for position, chart_data in enumerate(report.charts):
                summary.append([self._cell(summary, ""), self._cell(summary, "")])
                current_row += 1
                summary.append(
                    [
                        self._cell(summary, chart_data.title, bold=True),
                        self._cell(summary, "Count", bold=True),
                    ]
                )
                current_row += 1
                header_row = current_row
                for category, value in zip(chart_data.categories, chart_data.values, strict=True):
                    summary.append(
                        [self._cell(summary, category), self._cell(summary, value, align="right")]
                    )
                    current_row += 1
                chart = PieChart() if position == 0 else BarChart()
                if isinstance(chart, BarChart):
                    chart.type = "col"
                    chart.legend = None
                chart.style = 10 + position
                chart.title = chart_data.title
                chart.height = 7
                chart.width = 12
                values = Reference(
                    summary,
                    min_col=2,
                    min_row=header_row,
                    max_row=header_row + len(chart_data.values),
                )
                categories = Reference(
                    summary,
                    min_col=1,
                    min_row=header_row + 1,
                    max_row=header_row + len(chart_data.categories),
                )
                chart.add_data(values, titles_from_data=True)
                chart.set_categories(categories)
                summary.add_chart(chart, f"D{4 + position * 15}")

            data = workbook.create_sheet(report.data_sheet_title[:31])
            data.sheet_properties.tabColor = self._teal
            data.sheet_view.showGridLines = False
            data.freeze_panes = "A2"
            headers = report.dataframe.columns
            for index, header in enumerate(headers, start=1):
                data.column_dimensions[get_column_letter(index)].width = min(
                    42, max(14, len(str(header)) + 4)
                )
            data.append(
                [
                    self._cell(
                        data,
                        header,
                        bold=True,
                        color=self._white,
                        fill=self._navy,
                        align="center",
                    )
                    for header in headers
                ]
            )
            for row_number, row in enumerate(report.dataframe.iter_rows(), start=2):
                fill = self._pale_blue if row_number % 2 == 0 else "E8F5E9"
                data.append([self._cell(data, value, fill=fill) for value in row])
            if headers:
                data.auto_filter.ref = (
                    f"A1:{get_column_letter(len(headers))}{report.dataframe.height + 1}"
                )
            workbook.save(destination)
        finally:
            workbook.close()


@dataclass(frozen=True, slots=True)
class ReportLookupPolicy:
    """Explicit trusted-code opt-in for a bounded AI projection of a registered query."""

    permission: str
    fields: tuple[str, ...]
    classification: str
    max_rows: int = 10

    def __post_init__(self) -> None:
        if not self.permission or not self.fields or len(set(self.fields)) != len(self.fields):
            raise ValueError("Lookup permission and unique projection fields are required")
        if self.classification not in {"PUBLIC", "INTERNAL", "CONFIDENTIAL"}:
            raise ValueError("Lookup data classification is invalid")
        if not 1 <= self.max_rows <= 100:
            raise ValueError("Lookup row limit is invalid")


@dataclass(frozen=True, slots=True)
class ReportDefinition:
    """Trusted server-side wiring and execution limits for one report kind."""

    key: str
    version: int
    processor: type[BasePolarsReport]
    renderer: type[BaseExcelRenderer]
    chunk_size: int = 100
    max_rows: int = 100_000
    parallelism: int = 1
    lifetime_days: int = 10
    priority: int = 5
    ai_lookup: ReportLookupPolicy | None = None
