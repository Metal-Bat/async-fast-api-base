"""Tests for report workbook rendering and encrypted archives."""

from datetime import UTC, datetime
from pathlib import Path

import polars as pl
import pyzipper
from openpyxl import load_workbook
from openpyxl.chart import PieChart

from apps.reporting.application.base import CleanExcelRenderer
from apps.reporting.tasks import _create_archive
from apps.users.application.reporting import UserPolarsReport


def source_frame() -> pl.DataFrame:
    return pl.from_dicts(
        [
            {
                "username": "admin",
                "email": "admin@example.com",
                "first_name": "Ada",
                "last_name": "Admin",
                "is_superuser": True,
                "is_active": True,
                "created_at": datetime(2026, 1, 1, tzinfo=UTC),
                "updated_at": None,
                "deleted_at": None,
            },
            {
                "username": "archived",
                "email": None,
                "first_name": None,
                "last_name": None,
                "is_superuser": False,
                "is_active": False,
                "created_at": datetime(2026, 1, 2, tzinfo=UTC),
                "updated_at": datetime(2026, 2, 1, tzinfo=UTC),
                "deleted_at": datetime(2026, 2, 2, tzinfo=UTC),
            },
        ]
    )


def test_user_report_builds_clean_streaming_workbook_and_aes_zip(tmp_path: Path) -> None:
    processor = UserPolarsReport()
    chunk = processor.transform_chunk(source_frame(), "en")
    report = processor.finalize([chunk], "en")
    xlsx_path = tmp_path / "users.xlsx"
    zip_path = tmp_path / "users.zip"

    CleanExcelRenderer().render(report, xlsx_path)
    workbook = load_workbook(xlsx_path, read_only=False, data_only=True)
    try:
        assert workbook.sheetnames == ["Summary", "Users"]
        assert workbook["Users"].max_row == 3
        assert len(workbook["Summary"]._charts) == 3
        assert isinstance(workbook["Summary"]._charts[0], PieChart)
        assert workbook["Summary"].sheet_properties.tabColor is not None
    finally:
        workbook.close()

    _create_archive(xlsx_path, zip_path, "strong-password")
    with pyzipper.AESZipFile(zip_path) as archive:
        archive.setpassword(b"strong-password")
        assert archive.read("users.xlsx")[:2] == b"PK"
